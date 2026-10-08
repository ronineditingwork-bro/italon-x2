#!/usr/bin/env python3
"""Импорт фото артикулов из папок с JPG/PNG/WebP: имя файла = артикул (например 610010001530.jpg).

Использование:
  python3 scripts/import-photos-folder.py --src папка1 [--src папка2 ...] [--sources таблица.xlsx] [--skip артикул ...]

Файлы переводятся в WebP (до 900 px по большей стороне), кладутся в public/media/products,
в data/product-images.json добавляются записи kind=sku, в data/image-provenance.json — происхождение.
Нужен Pillow; --sources (openpyxl) необязателен: из таблицы берутся ссылки на страницу-источник.
"""
import argparse, hashlib, io, json
from pathlib import Path
from PIL import Image, ImageOps

ap = argparse.ArgumentParser()
ap.add_argument('--src', action='append', required=True)
ap.add_argument('--sources')
ap.add_argument('--skip', nargs='*', default=[])
ap.add_argument('--origin', default='Файл загружен владельцем сайта в репозиторий; имя файла равно артикулу')
args = ap.parse_args()
root = Path(__file__).resolve().parents[1]
price = json.loads((root / 'data/italon-price-2026-07.json').read_text())
targets = {p['code'] for s in price['series'] for p in s['items']}
manifest = json.loads((root / 'data/product-images.json').read_text())
provenance = json.loads((root / 'data/image-provenance.json').read_text())
sources = {}
if args.sources:
    from openpyxl import load_workbook
    for row in load_workbook(args.sources, read_only=True).active.iter_rows(min_row=2, values_only=True):
        if row[2]: sources[str(row[2])] = {'imageUrl': row[9], 'sourceUrl': row[11]}
dest = root / 'public/media/products'; dest.mkdir(parents=True, exist_ok=True)
done, problems = 0, []
for folder in args.src:
    for path in sorted(p for p in Path(folder).iterdir() if p.suffix.lower() in ('.jpg', '.jpeg', '.png', '.webp')):
        code = path.stem
        if code in args.skip: problems.append((code, 'пропущен по просьбе')); continue
        if code not in targets: problems.append((code, 'нет в прайсе')); continue
        raw = path.read_bytes()
        try:
            with Image.open(io.BytesIO(raw)) as loaded:
                image = ImageOps.exif_transpose(loaded).convert('RGB')
        except Exception as error:
            problems.append((code, f'не открывается: {error}')); continue
        if min(image.size) < 100: problems.append((code, f'слишком мелкое: {image.size}')); continue
        image.thumbnail((900, 900), Image.Resampling.LANCZOS)
        encoded = io.BytesIO(); image.save(encoded, format='WEBP', quality=76, method=6)
        data = encoded.getvalue(); name = hashlib.sha256(data).hexdigest()[:20] + '.webp'
        (dest / name).write_bytes(data)
        meta = {'src': '/media/products/' + name, 'kind': 'sku', 'width': image.width, 'height': image.height}
        manifest['products'][code] = meta
        provenance['products'][code] = {'code': code, **sources.get(code, {}), 'evidence': args.origin,
                                        'originalSha256': hashlib.sha256(raw).hexdigest(), 'asset': meta['src']}
        done += 1
(root / 'data/product-images.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n')
(root / 'data/image-provenance.json').write_text(json.dumps(provenance, ensure_ascii=False, indent=2) + '\n')
used = {p['src'].rsplit('/', 1)[1] for group in manifest.values() for p in group.values()}
for file in dest.glob('*.webp'):
    if file.name not in used: file.unlink()
print(json.dumps({'imported': done, 'problems': problems, 'skuImages': len(manifest['products']),
                  'bytes': sum((dest / f).stat().st_size for f in used)}, ensure_ascii=False))
